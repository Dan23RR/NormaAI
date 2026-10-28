"""Find the legal references in a document: acts, provisions, quotes.

Pure functions, no network. The patterns cover how Italian and English
documents usually cite EU and Italian law:

- EU acts: "Regolamento (UE) 2016/679", "Regulation (EU) 2024/1689",
  "Direttiva 2013/34/UE", "Regolamento (CE) n. 1907/2006", and the short names
  (GDPR, AI Act, NIS2, DORA, CSRD, CSDDD, CRA, ...).
- Names a document defines: "Regolamento (UE) 2016/679 (di seguito
  «Regolamento»)" makes "Regolamento" mean that act from there on.
- Italian acts: "d.lgs. 196/2003", "decreto legislativo 30 giugno 2003, n. 196",
  "legge 132/2025", "Codice privacy". NormaAI does not hold Italian law yet, so
  they are reported as references to check on Normattiva.
- Provisions: "art. 13", "artt. 13 e 14", "articolo 50, paragrafo 5",
  "Art. 6(1)(f)", "art. 4 bis", "considerando 47", "Allegato III".

A reference takes the act written right after it ("art. 13 del GDPR") or right
before it ("GDPR, art. 13"); failing that, the last act named earlier in the
same paragraph, marked as inferred; failing that, the only EU act the document
names. A quote («...», “...”, "...") belongs to the nearest reference in its
paragraph.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from src.agents.citation_verifier import parse_article

if TYPE_CHECKING:
    from pathlib import Path

# ─── Acts ────────────────────────────────────────────────────────────

# Short names of the acts NormaAI tracks (and SFDR, which CELLAR serves too).
ACT_NAMES: dict[str, str] = {
    "32016R0679": "GDPR",
    "32024R1689": "AI Act",
    "32022L2555": "NIS2",
    "32022R2554": "DORA",
    "32022L2464": "CSRD",
    "32024L1760": "CSDDD",
    "32024R2847": "Cyber Resilience Act",
    "32020R0852": "Tassonomia",
    "32023R2772": "ESRS",
    "32013L0034": "Direttiva contabile",
    "32019R2088": "SFDR",
    "32026R1744": "Omnibus digitale sull'IA",
    "32026L0470": "Omnibus I",
}

_SECTOR = {
    "regolamento": "R",
    "regulation": "R",
    "direttiva": "L",
    "directive": "L",
    "decisione": "D",
    "decision": "D",
}
_ORG = r"(?:UE|EU|CE|EC|CEE|EEC)"
_EU_ACT_RE = re.compile(
    r"\b(?P<kind>regolamento|direttiva|decisione|regulation|directive|decision)"
    r"(?:\s+(?:delegato|di\s+esecuzione|delegated|implementing))?"
    rf"\s*\(\s*{_ORG}\s*\)\s*"
    r"(?P<n>n\.\s*|no\.?\s*)?(?P<a>\d{1,4})\s*/\s*(?P<b>\d{1,4})(?!\d)",
    re.IGNORECASE,
)
# The older form of directives: "Direttiva 2013/34/UE".
_EU_DIRECTIVE_RE = re.compile(
    rf"\b(?P<kind>direttiva|directive)\s+(?P<a>\d{{4}})/(?P<b>\d{{1,4}})/{_ORG}\b",
    re.IGNORECASE,
)
_SHORT_NAMES: tuple[tuple[re.Pattern[str], str], ...] = (
    (re.compile(r"\b(?:GDPR|RGPD)\b"), "32016R0679"),
    (
        re.compile(
            r"\bAI\s?Act\b|\b[Rr]egolamento\s+(?:europeo\s+)?(?:sull['’]\s?)?"
            r"(?:IA|intelligenza\s+artificiale)\b"
        ),
        "32024R1689",
    ),
    (re.compile(r"\bNIS\s?2\b"), "32022L2555"),
    (re.compile(r"\bDORA\b"), "32022R2554"),
    (re.compile(r"\bCSRD\b"), "32022L2464"),
    (re.compile(r"\b(?:CSDDD|CS3D)\b"), "32024L1760"),
    (re.compile(r"\bCRA\b|\bCyber\s+Resilience\s+Act\b"), "32024R2847"),
    (
        re.compile(r"\b(?:[Rr]egolamento\s+)?[Tt]assonomia\b|\bTaxonomy\s+Regulation\b"),
        "32020R0852",
    ),
    (re.compile(r"\bESRS\b"), "32023R2772"),
    (re.compile(r"\bSFDR\b"), "32019R2088"),
)

_IT_MONTHS = (
    "gennaio|febbraio|marzo|aprile|maggio|giugno|luglio|agosto|settembre|ottobre|novembre|dicembre"
)
_IT_ACT_RE = re.compile(
    r"\b(?P<kind>d\.\s?lgs\.?|decreto\s+legislativo|legge|l\.)\s*"
    r"(?:(?:n\.\s*)?(?P<n1>\d{1,4})\s*/\s*(?P<y1>\d{4})"
    rf"|(?:\d{{1,2}}\s+(?:{_IT_MONTHS})\s+(?P<y2>\d{{4}})\s*,?\s*)?n\.\s*(?P<n2>\d{{1,4}})"
    r"(?:\s+del\s+(?P<y3>\d{4}))?)",
    re.IGNORECASE,
)
_IT_NAMES: tuple[tuple[re.Pattern[str], str], ...] = (
    (
        re.compile(
            r"\b[Cc]odice\s+(?:della\s+)?privacy\b|\b[Cc]odice\s+in\s+materia\s+di\s+protezione"
            r"\s+dei\s+dati\s+personali\b"
        ),
        "IT:dlgs:2003:196",
    ),
    (re.compile(r"\b[Ss]tatuto\s+dei\s+lavoratori\b"), "IT:legge:1970:300"),
)
# "(di seguito «Regolamento» o «GDPR»)" right after an act: names the document defines.
_DEFINITION_RE = re.compile(r"\s*\((?P<body>[^()]{0,200})\)")
_DEFINED_NAME_RE = re.compile(r"[«“\"]\s*([^«»“”\"]{2,40}?)\s*[»”\"]")


@dataclass(frozen=True)
class Act:
    """An act a reference belongs to."""

    key: str  # CELEX of an EU act ("32024R1689"); "IT:dlgs:2003:196" for Italian law
    label: str  # how the report names it

    @property
    def eu(self) -> bool:
        return not self.key.startswith("IT:")

    @property
    def normattiva_url(self) -> str | None:
        """Normattiva page of an Italian act (URN with year and number)."""
        if self.eu:
            return None
        _, kind, year, number = self.key.split(":")
        urn_kind = "decreto.legislativo" if kind == "dlgs" else "legge"
        return f"https://www.normattiva.it/uri-res/N2Ls?urn:nir:stato:{urn_kind}:{year};{number}"


@dataclass(frozen=True)
class _Mention:
    act: Act
    start: int
    end: int


def eu_act(celex: str, written: str | None = None) -> Act:
    """The Act for a CELEX, labelled with its short name when NormaAI knows it."""
    name = ACT_NAMES.get(celex)
    written = re.sub(r"\s+", " ", written or "").strip()
    if written and name and name.lower() not in written.lower():
        return Act(celex, f"{written} ({name})")
    return Act(celex, written or name or celex)


def italian_act(kind: str, year: str, number: str) -> Act:
    """The Act for an Italian decree or law."""
    label = f"D.Lgs. {number}/{year}" if kind == "dlgs" else f"Legge {number}/{year}"
    if (kind, year, number) == ("dlgs", "2003", "196"):
        label += " (Codice privacy)"
    return Act(f"IT:{kind}:{year}:{int(number)}", label)


def _eu_celex(kind: str, a: str, b: str, number_first: bool) -> str | None:
    """CELEX of "Regolamento (UE) 2016/679" (year/number) or "(CE) n. 1907/2006"."""
    sector = _SECTOR[kind.lower()]
    if number_first:
        number, year = a, b
    elif len(a) == 4 and 1950 <= int(a) <= 2100:
        year, number = a, b
    elif len(b) == 4 and 1950 <= int(b) <= 2100:
        number, year = a, b
    else:
        return None
    return f"3{year}{sector}{int(number):04d}"


def _act_mentions(text: str) -> list[_Mention]:
    mentions: list[_Mention] = []
    for m in _EU_ACT_RE.finditer(text):
        celex = _eu_celex(m.group("kind"), m.group("a"), m.group("b"), bool(m.group("n")))
        if celex:
            mentions.append(_Mention(eu_act(celex, m.group(0)), m.start(), m.end()))
    for m in _EU_DIRECTIVE_RE.finditer(text):
        celex = _eu_celex(m.group("kind"), m.group("a"), m.group("b"), number_first=False)
        if celex:
            mentions.append(_Mention(eu_act(celex, m.group(0)), m.start(), m.end()))
    for m in _IT_ACT_RE.finditer(text):
        kind = "legge" if m.group("kind").lower().startswith("l") else "dlgs"
        year = m.group("y1") or m.group("y2") or m.group("y3")
        number = m.group("n1") or m.group("n2")
        if year and number:
            mentions.append(_Mention(italian_act(kind, year, number), m.start(), m.end()))
    for pattern, key in _IT_NAMES:
        _, kind, year, number = key.split(":")
        for m in pattern.finditer(text):
            mentions.append(_Mention(italian_act(kind, year, number), m.start(), m.end()))
    for pattern, celex in _SHORT_NAMES:
        for m in pattern.finditer(text):
            mentions.append(_Mention(eu_act(celex), m.start(), m.end()))
    mentions = _drop_overlaps(mentions)
    return _drop_overlaps(mentions + _defined_names(text, mentions))


def _drop_overlaps(mentions: list[_Mention]) -> list[_Mention]:
    """Keep the longest mention where two overlap ("Regolamento (UE) 2016/679" over "GDPR")."""
    kept: list[_Mention] = []
    for m in sorted(mentions, key=lambda x: (x.start, -(x.end - x.start))):
        if kept and m.start < kept[-1].end:
            continue
        kept.append(m)
    return kept


def _defined_names(text: str, mentions: list[_Mention]) -> list[_Mention]:
    """Mentions of the names the document defines for an act, after the definition."""
    found: list[_Mention] = []
    for mention in mentions:
        d = _DEFINITION_RE.match(text, mention.end)
        if not d:
            continue
        for name in _DEFINED_NAME_RE.findall(d.group("body")):
            name = name.strip()
            if not name or not name[0].isupper():
                continue
            pattern = re.compile(rf"(?<![\w«“\"]){re.escape(name)}(?![\w»”\"])")
            for m in pattern.finditer(text, d.end()):
                found.append(_Mention(mention.act, m.start(), m.end()))
    return found


# ─── Provisions ──────────────────────────────────────────────────────

_LATIN = "bis|ter|quater|quinquies|sexies|septies|octies|novies|nonies|decies"
# A number that is not the first half of an act number ("2016/679").
_NUM = rf"\d{{1,4}}(?!\d|\s*/)(?:[\s-]?(?:{_LATIN})\b|[a-z](?![a-z]))?"
_SUBDIV = r"(?:\s*\(\s*\d{1,3}\s*\))?(?:\s*\(\s*[a-z]{1,2}\s*\))*"
_PARA = (
    r"(?:\s*,?\s*(?:paragrafi|paragrafo|parr\.|par\.|paragraphs|paragraph|para\.|commi|comma|co\.)"
    r"\s*\d{1,3}(?:\s*(?:,|e|ed|and|-|–)\s*\d{1,3})*)?"
)
_POINT = r"(?:\s*,?\s*(?:lettera|lett\.|letter|point|punto)\s*[a-z0-9]{1,3}\)?)?"
_ITEM = rf"{_NUM}{_SUBDIV}{_PARA}{_POINT}"
_SEP = r"\s*(?:,|\be\b|\bed\b|\band\b|-|–)\s*"
_ARTICLE_RE = re.compile(
    rf"\b(?:artt?\.|art\b|articol[oi]|articles?)\s*(?P<items>{_ITEM}(?:{_SEP}{_ITEM})*)",
    re.IGNORECASE,
)
_ITEM_RE = re.compile(rf"(?P<num>{_NUM})(?P<sub>{_SUBDIV})(?P<para>{_PARA}){_POINT}", re.IGNORECASE)
# "paragrafo 9 dell'articolo 50", "comma 2 dell'art. 13": the paragraph first.
_PARAGRAPH_FIRST_RE = re.compile(
    r"\b(?:paragrafo|par\.|comma|co\.)\s*(?P<para>\d{1,3})\s+dell['’]\s*"
    rf"(?:articolo|art\.?)\s*(?P<num>{_NUM})",
    re.IGNORECASE,
)
_RECITAL_RE = re.compile(
    r"\b(?:considerand[oi]|recitals?)\s*(?:n\.\s*)?"
    r"(?P<items>\d{1,3}(?:\s*(?:,|\be\b|\band\b|-|–)\s*\d{1,3})*)",
    re.IGNORECASE,
)
_ANNEX_RE = re.compile(r"\b(?:[Aa]llegat[oi]|[Aa]nnex(?:es)?)\s+(?P<num>[IVXLC]{1,6}|\d{1,2})\b")


@dataclass(frozen=True)
class Provision:
    """One provision a reference names: an article, a recital or an annex."""

    key: str  # parse_article key: "art:13", "art:4bis", "recital:47", "annex:III"
    paragraphs: tuple[str, ...] = ()

    @property
    def label(self) -> str:
        kind, _, number = self.key.partition(":")
        if kind == "recital":
            return f"considerando {number}"
        if kind == "annex":
            return f"allegato {number}"
        m = re.match(rf"^(\d+)({_LATIN})$", number)
        base = f"art. {m.group(1)} {m.group(2)}" if m else f"art. {number}"
        if self.paragraphs:
            word = "par." if len(self.paragraphs) == 1 else "parr."
            return f"{base}, {word} {', '.join(self.paragraphs)}"
        return base


@dataclass
class Reference:
    """A reference to one or more provisions of an act, as written in the document."""

    start: int
    end: int
    provisions: list[Provision]
    act: Act | None = None
    act_inferred: bool = False
    quotes: list[str] = field(default_factory=list)
    raw: str = ""  # the reference with its act, as written: "art. 13 del GDPR"
    line: int = 0  # 1-based line (paragraph) of the document
    context: str = ""  # the text around the reference


def _numbers(text: str) -> list[str]:
    """Paragraph numbers in "par. 1", "commi 1-3", "(2)"; ranges expanded."""
    found: list[str] = []
    for m in re.finditer(r"(\d{1,3})(?:\s*[-–]\s*(\d{1,3}))?", text):
        first, last = int(m.group(1)), int(m.group(2) or m.group(1))
        if last < first or last - first > 20:
            last = first
        found.extend(str(n) for n in range(first, last + 1))
    return found


def _article_provisions(items: str) -> list[Provision]:
    provisions: list[Provision] = []
    previous_end = 0
    for m in _ITEM_RE.finditer(items):
        sub = re.match(r"\s*\(\s*(\d{1,3})\s*\)", m.group("sub"))
        paragraphs = [sub.group(1)] if sub else _numbers(m.group("para"))
        key = parse_article(f"Art. {m.group('num')}")
        if key is None:
            continue
        # "artt. 15-22": a range between two plain article numbers.
        between = items[previous_end : m.start()]
        if provisions and re.search(r"[-–]", between):
            first = re.fullmatch(r"art:(\d+)", provisions[-1].key)
            last = re.fullmatch(r"art:(\d+)", key)
            if first and last and 0 < int(last.group(1)) - int(first.group(1)) <= 30:
                for n in range(int(first.group(1)) + 1, int(last.group(1))):
                    provisions.append(Provision(f"art:{n}"))
        provisions.append(Provision(key, tuple(dict.fromkeys(paragraphs))))
        previous_end = m.end()
    return provisions


def _references(text: str) -> list[Reference]:
    refs: list[Reference] = []
    for m in _PARAGRAPH_FIRST_RE.finditer(text):
        key = parse_article(f"Art. {m.group('num')}")
        if key is not None:
            refs.append(Reference(m.start(), m.end(), [Provision(key, (m.group("para"),))]))
    taken = [(r.start, r.end) for r in refs]
    for m in _ARTICLE_RE.finditer(text):
        if any(start <= m.start() < end for start, end in taken):
            continue
        provisions = _article_provisions(m.group("items"))
        if provisions:
            refs.append(Reference(m.start(), m.end(), provisions))
    for m in _RECITAL_RE.finditer(text):
        provisions = [Provision(f"recital:{n}") for n in dict.fromkeys(_numbers(m.group("items")))]
        refs.append(Reference(m.start(), m.end(), provisions))
    for m in _ANNEX_RE.finditer(text):
        refs.append(Reference(m.start(), m.end(), [Provision(f"annex:{m.group('num').upper()}")]))
    return sorted(refs, key=lambda r: r.start)


# ─── Linking references, acts and quotes ─────────────────────────────

_ABBREVIATIONS = re.compile(
    r"\b(?:artt?|parr?|n|lett|co|cfr|es|ecc|d\.lgs|lgs|l|reg|dir|no|para)\.", re.IGNORECASE
)
_QUOTE_RE = re.compile(
    r"«\s*(?P<a>[^«»\n]{20,1500}?)\s*»|“(?P<b>[^“”\n]{20,1500}?)”|\"(?P<c>[^\"\n]{20,1500}?)\""
)
_AFTER_WINDOW = 90
_BEFORE_WINDOW = 40


def _crosses_sentence(gap: str) -> bool:
    gap = _ABBREVIATIONS.sub(" ", gap)
    return bool(re.search(r"[\n;]|[.!?]\s+[A-ZÀ-Ý«“\"]", gap))


def _line_bounds(text: str, position: int) -> tuple[int, int]:
    start = text.rfind("\n", 0, position) + 1
    end = text.find("\n", position)
    return start, len(text) if end == -1 else end


def _context(text: str, start: int, end: int, width: int = 160) -> str:
    line_start, line_end = _line_bounds(text, start)
    left = max(line_start, start - width)
    right = min(line_end, end + width)
    snippet = text[left:right].strip()
    return ("… " if left > line_start else "") + snippet + (" …" if right < line_end else "")


def _link_acts(text: str, refs: list[Reference], mentions: list[_Mention]) -> None:
    eu_acts = {m.act.key for m in mentions if m.act.eu}
    only_act = next((m.act for m in mentions if m.act.eu), None) if len(eu_acts) == 1 else None
    for ref in refs:
        line_start, _ = _line_bounds(text, ref.start)
        after = next(
            (
                m
                for m in mentions
                if ref.end <= m.start <= ref.end + _AFTER_WINDOW
                and not _crosses_sentence(text[ref.end : m.start])
            ),
            None,
        )
        before = next(
            (
                m
                for m in reversed(mentions)
                if ref.start - _BEFORE_WINDOW <= m.end <= ref.start
                and re.fullmatch(r"[\s,(:\-–]*", text[m.end : ref.start])
            ),
            None,
        )
        earlier = next(
            (m for m in reversed(mentions) if line_start <= m.end <= ref.start),
            None,
        )
        # "Codice privacy, art. 2-ter, e la Direttiva ...": the closer act wins.
        if after and before and ref.start - before.end < after.start - ref.end:
            after = None
        if after:
            ref.act, ref.raw = after.act, text[ref.start : after.end]
        elif before:
            ref.act, ref.raw = before.act, text[before.start : ref.end]
        elif earlier:
            ref.act, ref.act_inferred = earlier.act, True
        elif only_act:
            ref.act, ref.act_inferred = only_act, True
        if not ref.raw:
            ref.raw = text[ref.start : ref.end]
        ref.raw = re.sub(r"\s+", " ", ref.raw).strip()


def _link_quotes(text: str, refs: list[Reference]) -> None:
    for q in _QUOTE_RE.finditer(text):
        quote = next(g for g in q.groups() if g)
        line = _line_bounds(text, q.start())
        best: tuple[int, Reference] | None = None
        for ref in refs:
            if _line_bounds(text, ref.start) != line:
                continue
            if ref.end <= q.start() and q.start() - ref.end <= 300:
                distance = q.start() - ref.end
            elif ref.start >= q.end() and ref.start - q.end() <= 150:
                distance = 2 * (ref.start - q.end())
            else:
                continue
            if best is None or distance < best[0]:
                best = (distance, ref)
        if best:
            best[1].quotes.append(quote.strip())


def extract_references(text: str) -> list[Reference]:
    """Every reference to a provision in ``text``, with its act and quotes."""
    mentions = _act_mentions(text)
    refs = [r for r in _references(text) if not any(m.start <= r.start < m.end for m in mentions)]
    _link_acts(text, refs, mentions)
    _link_quotes(text, refs)
    for ref in refs:
        ref.line = text.count("\n", 0, ref.start) + 1
        ref.context = _context(text, ref.start, ref.end)
    return refs


def mentioned_acts(text: str) -> list[Act]:
    """The acts a document names, in order of first mention."""
    seen: dict[str, Act] = {}
    for m in _act_mentions(text):
        seen.setdefault(m.act.key, m.act)
    return list(seen.values())


# ─── Documents ───────────────────────────────────────────────────────

_IT_WORDS = frozenset(
    "il lo la gli le di del della dei delle che per con non sono è ai sensi articolo dati".split()
)
_EN_WORDS = frozenset("the of and to in is are that for with by article data shall".split())


def detect_language(text: str) -> str:
    """Return "IT" or "EN", by counting common words (Italian when in doubt)."""
    words = re.findall(r"[a-zàèéìòù]+", text.lower())
    italian = sum(w in _IT_WORDS for w in words)
    english = sum(w in _EN_WORDS for w in words)
    return "EN" if english > italian else "IT"


def load_text(path: Path) -> str:
    """Plain text of a .txt, .md, .html, .docx or .pdf document."""
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md", ".text"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return re.sub(r"\n{3,}", "\n\n", soup.get_text("\n"))
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        from pypdf import PdfReader

        raw = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        # PDF lines break mid-sentence: rejoin hyphenated words and wrapped lines.
        raw = re.sub(r"(\w)-\n(\w)", r"\1\2", raw)
        return re.sub(r"(?<![.:;])\n(?!\n)", " ", raw)
    raise ValueError(f"Formato non supportato: {suffix or 'senza estensione'}")
